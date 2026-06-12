#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
RoboServe 大三实习面试备战指南
定位：大三实习生。重点：项目介绍 + 八股文 + 算法入门 + 简历
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ============================================================
style = doc.styles['Normal']
font = style.font
font.name = '\u5fae\u8f6f\u96c5\u9ed1'
font.size = Pt(10.5)
font.color.rgb = RGBColor(0x33, 0x33, 0x33)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '\u5fae\u8f6f\u96c5\u9ed1')

for level in range(1, 5):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = '\u5fae\u8f6f\u96c5\u9ed1'
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), '\u5fae\u8f6f\u96c5\u9ed1')
    if level == 1:
        hs.font.size = Pt(16)
        hs.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
    elif level == 2:
        hs.font.size = Pt(13)
        hs.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    elif level == 3:
        hs.font.size = Pt(11.5)
        hs.font.color.rgb = RGBColor(0x34, 0x49, 0x5E)

# ============================================================
def add_h1(text):
    doc.add_heading(text, level=1)

def add_h2(text):
    doc.add_heading(text, level=2)

def add_h3(text):
    doc.add_heading(text, level=3)

def add_p(text):
    p = doc.add_paragraph()
    p.add_run(text)
    return p

def add_code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.8)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x2D, 0x2D, 0x2D)
    run.element.rPr.rFonts.set(qn('w:eastAsia'), '\u5fae\u8f6f\u96c5\u9ed1')
    pPr = p._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F7FA')
    pPr.append(shd)
    return p

def add_qa(q, a):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run(q)
    run.bold = True
    run.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    run.font.size = Pt(10.5)
    p2 = doc.add_paragraph()
    p2.paragraph_format.left_indent = Cm(0.5)
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(a)
    r2.font.color.rgb = RGBColor(0x27, 0xAE, 0x60)

def add_warn(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    run = p.add_run('\u26a0\ufe0f ' + text)
    run.font.color.rgb = RGBColor(0xE6, 0x7E, 0x22)
    run.font.size = Pt(10)
    run.bold = True

def add_tip(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    run = p.add_run('\ud83d\udca1 ' + text)
    run.font.color.rgb = RGBColor(0xE6, 0x7E, 0x22)
    run.font.size = Pt(10)

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ''
            cell.paragraphs[0].add_run(str(val)).font.size = Pt(10)
    doc.add_paragraph('')

def add_page_break():
    doc.add_page_break()

def add_bold_p(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    return p

# ============================================================
# 封面
# ============================================================
for _ in range(4):
    doc.add_paragraph('')

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('\u5927\u4e09\u5b9e\u4e60\u9762\u8bd5\u5907\u6218\u6307\u5357'); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

t2 = doc.add_paragraph(); t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = t2.add_run('\u9879\u76ee\u4ecb\u7ecd \u00d7 \u516b\u80a1\u6587 \u00d7 \u7b97\u6cd5\u5165\u95e8 \u00d7 \u7b80\u5386'); r2.font.size = Pt(15); r2.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)

doc.add_paragraph('')
t3 = doc.add_paragraph(); t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
t3.add_run('\u5b9a\u4f4d\uff1a\u5927\u4e09\u5b9e\u4e60\u751f\uff08\u975e\u793e\u62db\uff09 | \u9879\u76ee\uff1aRoboServe \u667a\u80fd\u5ba2\u670d\u7cfb\u7edf').font.size = Pt(11)

doc.add_paragraph('')
t4 = doc.add_paragraph(); t4.alignment = WD_ALIGN_PARAGRAPH.CENTER
t4.add_run('2026\u5e746\u670812\u65e5').font.size = Pt(10)

add_page_break()

# ============================================================
# 第0章
# ============================================================
add_h1('\u7b2c\u96f6\u7ae0\uff1a\u5927\u4e09\u5b9e\u4e60\u9762\u8bd5\u662f\u4ec0\u4e48\u6837\u7684')

add_h2('\u9762\u8bd5\u5b98\u9762\u5bf9\u5927\u4e09\u5b9e\u4e60\u751f\uff0c\u53ea\u8003\u5bdf\u4e09\u4ef6\u4e8b')

add_table(
    ['\u9762\u8bd5\u5b98\u5fc3\u7406', '\u4f60\u8981\u600e\u4e48\u8bc1\u660e'],
    [
        ['\u2460 \u8fd9\u4e2a\u4eba\u80fd\u4e0d\u80fd\u5b66\u5f97\u52a8\uff1f', '\u7528\u9879\u76ee\u8bc1\u660e\u2014\u2014\u201c\u6211\u72ec\u7acb\u4ece0\u642d\u5efa\u4e86\u8fd9\u4e2a\u7cfb\u7edf\uff0c\u4ece\u67b6\u6784\u5230\u4ee3\u7801\u5230\u8c03\u8bd5\u201d'],
        ['\u2461 \u57fa\u7840\u624e\u4e0d\u624e\u5b9e\uff1f', '\u516b\u80a1\u6587\u80fd\u7b54\u4e0a\u6765\u2014\u2014\u201c\u8fd9\u4e2a\u6211\u5b66\u6821\u5b66\u8fc7\uff0c\u9879\u76ee\u91cc\u4e5f\u7528\u5230\u4e86\u201d'],
        ['\u2462 \u6c9f\u901a\u6709\u6ca1\u95ee\u9898\uff1f', '\u8bf4\u8bdd\u6e05\u695a\u3001\u6001\u5ea6\u79ef\u6781\u3001\u4e0d\u4f1a\u5c31\u8bf4\u201c\u8fd9\u4e2a\u6211\u8fd8\u6ca1\u5b66\u5230\uff0c\u4f46\u6211\u53ef\u4ee5\u8bd5\u7740\u5206\u6790\u4e00\u4e0b\u201d'],
    ]
)

add_h2('\u9762\u8bd5\u65f6\u95f4\u5206\u914d\uff08\u5927\u4e09\u5b9e\u4e60\u771f\u5b9e\u6bd4\u4f8b\uff09')

add_code('''\u81ea\u6211\u4ecb\u7ecd+\u9879\u76ee\u4ecb\u7ecd  \u2588\u2588\u2588\u2588\u2588\u2588\u2588\u2588\u2588\u2588\u2588\u2588 40%  \u2190 \u6700\u91cd\u8981\uff01\u51b3\u5b9a\u7b2c\u4e00\u5370\u8c61
\u516b\u80a1\u6587                 \u2588\u2588\u2588\u2588\u2588\u2588\u2588\u2588\u2588 30%  \u2190 \u5b9e\u4e60\u751f\u5fc5\u8003
\u6280\u672f\u57fa\u7840 + \u7b97\u6cd5          \u2588\u2588\u2588\u2588\u2588 15%  \u2190 \u53ef\u80fd\u8003 Easy~Medium
\u9879\u76ee\u6df1\u6316\uff08\u53ef\u80fd\u6ca1\u6709\uff09  \u2588\u2588 10%  \u2190 \u4e0d\u8981\u4e3b\u52a8\u5c55\u5f00\u4f60\u4e0d\u61c2\u7684
\u53cd\u95ee\u73af\u8282               \u2588 5%     \u2190 \u51c6\u5907 2-3 \u4e2a\u95ee\u9898''')

add_warn('\u4f60\u4e4b\u524d\u7684\u6587\u6863\u91cc\u6709\u5f88\u591a\u201c\u6280\u672f\u6df1\u6316\u9898\u201d\uff0c\u4f46\u5927\u4e09\u5b9e\u4e60\u9762\u8bd5\u51e0\u4e4e\u4e0d\u4f1a\u95ee\u5230\u3002\u4f60\u53ea\u9700\u8981\u80fd\u8bf4\u6e05\u695a\u9879\u76ee\u6d41\u7a0b\u5c31\u591f\u4e86\uff01\u8d8a\u6df1\u5165\u7684\u70b9\u53cd\u800c\u5bb9\u6613\u88ab\u8ffd\u95ee\u5230\u4f60\u4e0d\u61c2\u7684\u5730\u65b9\u3002')

add_h2('\u590d\u4e60\u4f18\u5148\u7ea7\u6392\u5e8f')

add_table(
    ['\u4f18\u5148\u7ea7', '\u5185\u5bb9', '\u65f6\u95f4\u5206\u914d', '\u9762\u8bd5\u9891\u7387'],
    [
        ['\U0001f534 P0', '\u9879\u76ee\u4ecb\u7ecd\uff082\u5206\u949f / 30\u79d2\u7248\uff09', '30%', '\u6781\u9ad8'],
        ['\U0001f534 P0', '\u516b\u80a1\u6587\uff08Python/\u64cd\u4f5c\u7cfb\u7edf/\u7f51\u7edc/\u6570\u636e\u5e93\uff09', '35%', '\u6781\u9ad8'],
        ['\U0001f7e1 P1', 'LeetCode \u7b97\u6cd5\uff08Easy~Medium\uff09', '20%', '\u9ad8'],
        ['\U0001f7e1 P1', '\u9879\u76ee\u6280\u672f\u70b9\uff08\u5411\u91cf\u68c0\u7d22/\u6df7\u5408\u68c0\u7d22\u94fe\u8def/RRF\uff09', '10%', '\u4e2d'],
        ['\U0001f7e2 P2', '\u6df1\u5ea6\u6280\u672f\uff08A2A\u8bc4\u4f30/Memory\u6458\u8981\u8d28\u91cf\uff09', '5%', '\u4f4e'],
    ]
)

add_page_break()

# ============================================================
add_h1('\u7b2c\u4e00\u90e8\u5206\uff1a\u81ea\u6211\u4ecb\u7ecd + \u9879\u76ee\u4ecb\u7ecd\uff08\u9762\u8bd5\u7684 40%\uff09')

add_p('\u8fd9\u662f\u9762\u8bd5\u7b2c\u4e00\u95ee\uff0c\u76f4\u63a5\u51b3\u5b9a\u9762\u8bd5\u5b98\u5bf9\u4f60\u7684\u6574\u4f53\u5370\u8c61\u3002\u51c6\u5907\u4e09\u4e2a\u7248\u672c\uff0c\u6839\u636e\u9762\u8bd5\u5b98\u7684\u95ee\u6cd5\u5207\u6362\u3002')

add_h2('\u7248\u672c 1\uff1a2 \u5206\u949f\u5b8c\u6574\u7248\uff08\u201c\u4f60\u5148\u81ea\u6211\u4ecb\u7ecd\u4e00\u4e0b\u201d\uff09')

add_p('\u6309\u65f6\u95f4\u7ebf\u62c6\u89e3\uff1a')

add_code('''0~20\u79d2    \u6211\u662fXX\u5b66\u6821XX\u4e13\u4e1a\u5927\u4e09\u5b66\u751f\uff0c\u6700\u8fd1\u72ec\u7acb\u5f00\u53d1\u4e86\u4e00\u4e2a\u667a\u80fd\u5ba2\u670d\u7cfb\u7edf RoboServe\u3002

20~60\u79d2   \u7cfb\u7edf\u505a\u4ec0\u4e48\uff1a\u9488\u5bf9\u626b\u5730\u673a\u5668\u4eba\u7684\u6545\u969c\u6392\u67e5\u3001\u4f7f\u7528\u54a8\u8be2\u3001\u5929\u6c14\u67e5\u8be2\u3001\u62a5\u544a\u751f\u6210\u3002
         \u6280\u672f\u6808\uff1aPython + LangChain ReAct Agent + ChromaDB + Redis + Streamlit\u3002

60~100\u79d2  \u4e09\u4e2a\u6838\u5fc3\u4eae\u70b9\uff1a
         \u2460 \u6df7\u5408\u68c0\u7d22\u2014\u2014Dense\u5411\u91cf + BM25\u5173\u952e\u8bcd + RRF\u878d\u5408 + Cross-Encoder\u7cbe\u6392
         \u2461 Memory\u5206\u5c42\u7ba1\u7406\u2014\u2014\u89e3\u51b3\u957f\u5bf9\u8bdd token \u7206\u70b8\u95ee\u9898
         \u2462 A2A\u8bc4\u4f30Agent\u2014\u2014\u72ec\u7acb\u6253\u5206\u3001\u4e0d\u8fbe\u6807\u89e6\u53d1\u91cd\u751f\u6210

100~120\u79d2 \u5168\u94fe\u8def\u5df2\u8dd1\u901a\uff0c\u77e5\u8bc6\u5e93\u6709 119 \u6761\u6587\u6863\u3002
         LLM \u63a5\u7684\u901a\u4e49\u5343\u95ee\uff0c\u8bc4\u4f30 Agent \u7528\u7684\u7845\u57fa\u6d41\u52a8\u7684\u514d\u8d39\u6a21\u578b\u3002''')

add_h2('\u7248\u672c 2\uff1a30 \u79d2\u7cbe\u7b80\u7248\uff08\u201c\u7b80\u5355\u4ecb\u7ecd\u4e00\u4e0b\u9879\u76ee\u201d\uff09')

add_code('''\u6211\u505a\u4e86\u4e00\u4e2a\u667a\u80fd\u5ba2\u670d\u7cfb\u7edf RoboServe\uff0c\u57fa\u4e8e LangChain ReAct Agent\u3002
\u68c0\u7d22\u65b9\u9762\u7528\u4e86\u6df7\u5408\u68c0\u7d22\u2014\u2014Dense\u5411\u91cf + BM25 + RRF\u878d\u5408 + Cross-Encoder\u7cbe\u6392\uff0c\u56db\u9636\u6bb5\u4fdd\u969c\u53ec\u56de\u8d28\u91cf\u3002
\u8fd8\u505a\u4e86 Memory \u